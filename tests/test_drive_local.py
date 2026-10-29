"""Isolated integration tests; never require a running server or real user data."""
import io
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
from app.main import app
from app.database.database import Base
from app.database.session import get_db


@pytest.fixture
def local_client(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    (tmp_path / "uploads").mkdir()
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    sessions = sessionmaker(bind=engine)
    def database():
        with sessions() as session:
            yield session
    app.dependency_overrides[get_db] = database
    with TestClient(app) as client:
        yield client
    app.dependency_overrides.clear()
    engine.dispose()


def account(client, email):
    response = client.post("/auth/register", json={"full_name": "Test User", "email": email, "password": "strong-test-pass"})
    assert response.status_code == 200, response.text
    response = client.post("/auth/login", data={"username": email.upper(), "password": "strong-test-pass"})
    assert response.status_code == 200, response.text
    return {"Authorization": "Bearer " + response.json()["access_token"]}


def test_private_previews_ranges_archives_and_fallback(local_client):
    import zipfile
    c = local_client
    owner = account(c, "preview@example.com")
    other = account(c, "outsider@example.com")
    def upload(name, raw, mime="application/octet-stream"):
        r = c.post("/files/upload", headers=owner, files={"file": (name, raw, mime)})
        assert r.status_code == 200, r.text
        return r.json()["id"]
    fid = upload("movie.mp4", b"0123456789")
    r = c.get(f"/preview/{fid}/info", headers=owner)
    assert r.status_code == 200 and r.json()["kind"] == "video"
    token = r.json()["token"]
    assert c.get(f"/preview/{fid}/info", headers=other).status_code == 404
    assert c.get(f"/preview/{fid}/content?token=bad").status_code == 401
    streamed = c.get(f"/preview/{fid}/content?token={token}", headers={"Range": "bytes=2-5"})
    assert streamed.status_code == 206 and streamed.content == b"2345"
    assert streamed.headers["content-range"] == "bytes 2-5/10"
    assert c.get('/me', headers={"Authorization": "Bearer " + token}).status_code == 401
    fid2 = upload("other.mp4", b"private")
    assert c.get(f"/preview/{fid2}/content?token={token}").status_code == 401
    binary = upload("custom.weird", b"\x00\x01\xffABC")
    assert "41 42 43" in c.get(f"/preview/{binary}/info", headers=owner).json()["hex"]
    text = upload("source.html", b"<script>alert('unsafe')</script>")
    info = c.get(f"/preview/{text}/info", headers=owner).json()
    assert info["kind"] == "text" and "<script>" in info["text"]
    response = c.get(f'/preview/{text}/content?token={info["token"]}')
    assert response.headers["content-type"].startswith("text/plain")
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr("../readme.txt", "safe preview without filesystem extraction")
        z.writestr("bomb.txt", b"0" * 2000000)
    zipped = upload("bundle.zip", archive.getvalue())
    info = c.get(f"/preview/{zipped}/info", headers=owner).json()
    assert len(info["entries"]) == 2
    assert c.get(f"/preview/{zipped}/entry/0/info", headers=owner).json()["text"].startswith("safe preview")
    assert c.get(f"/preview/{zipped}/entry/1/info", headers=owner).status_code == 413
    assert c.get(f"/preview/{zipped}/entry/0/info", headers=other).status_code == 404
    assert c.get(f"/preview/{zipped}/entry/999/info", headers=owner).status_code == 404
    assert c.post(f"/drive/file/{fid}/trash", headers=owner, json={}).status_code == 200
    assert c.get(f"/preview/{fid}/content?token={token}").status_code == 404


def test_office_reading_views_escape_content(local_client):
    import zipfile
    from docx import Document
    from openpyxl import Workbook
    c = local_client
    headers = account(c, "officepreview@example.com")
    doc = Document()
    doc.add_paragraph("<script>private document</script>")
    word = io.BytesIO()
    doc.save(word)
    book = Workbook()
    book.active.append(["<script>private cell</script>", 12])
    excel = io.BytesIO()
    book.save(excel)
    slides = io.BytesIO()
    with zipfile.ZipFile(slides, "w") as z:
        z.writestr("ppt/slides/slide1.xml", '<a:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:t>Private slide</a:t></a:sld>')
    for name, raw, expected in [("report.docx", word.getvalue(), "private document"), ("sheet.xlsx", excel.getvalue(), "private cell"), ("slides.pptx", slides.getvalue(), "Private slide")]:
        r = c.post("/files/upload", headers=headers, files={"file": (name, raw, "application/octet-stream")})
        fid = r.json()["id"]
        info = c.get(f"/preview/{fid}/info", headers=headers).json()
        assert info["kind"] == "document"
        preview = c.get(f'/preview/{fid}/content?token={info["token"]}')
        assert preview.status_code == 200, preview.text
        assert expected in preview.text and "<script>" not in preview.text
        assert "sandbox" in preview.headers["content-security-policy"]


def test_login_validation_and_isolation(local_client):
    c = local_client
    assert c.get("/drive/items").status_code == 401
    assert c.post("/auth/register", json={"full_name": "A", "email": "a@example.com", "password": "short"}).status_code == 422
    owner = account(c, "owner@example.com")
    other = account(c, "other@example.com")
    assert c.post("/auth/login", data={"username": "owner@example.com", "password": "wrong"}).status_code == 401
    folder = c.post("/drive/folders", headers=owner, json={"name": "Private"}).json()["id"]
    assert c.post("/drive/folders", headers=other, json={"name": "Intruder", "parent_id": folder}).status_code == 404
    assert c.post(f"/drive/folder/{folder}/move", headers=other, json={}).status_code == 404
    assert c.get("/drive/items", headers=other).json()["items"] == []


def test_nested_copy_move_upload_trash_and_share(local_client):
    c = local_client
    headers = account(c, "drive@example.com")
    def post(url, data):
        r = c.post(url, headers=headers, json=data)
        assert r.status_code == 200, r.text
        return r.json()
    root = post("/drive/folders", {"name": "Projects"})["id"]
    child = post("/drive/folders", {"name": "Design", "parent_id": root})["id"]
    response = c.post("/files/upload", headers=headers, data={"folder_id": child}, files={"file": ("notes.txt", io.BytesIO(b"original bytes"), "text/plain")})
    assert response.status_code == 200, response.text
    file_id = response.json()["id"]
    assert response.json()["folder_id"] == child
    assert c.post(f"/drive/folder/{root}/move", headers=headers, json={"parent_id": child}).status_code == 400
    post(f"/drive/folder/{root}/copy", {})
    rows = c.get("/drive/items?view=home", headers=headers).json()["items"]
    cloned_file = next(i for i in rows if i["kind"] == "file" and i["id"] != file_id)
    assert c.get(f'/files/{cloned_file["id"]}/download', headers=headers).content == b"original bytes"
    post(f"/drive/file/{file_id}/star", {"starred": True})
    assert len(c.get("/drive/items?view=starred", headers=headers).json()["items"]) == 1
    post(f"/drive/file/{file_id}/rename", {"name": "renamed.txt"})
    assert len(c.get("/drive/items?q=renamed", headers=headers).json()["items"]) == 1
    share = post(f"/share/files/{file_id}", {"permission": "download"})
    assert c.get(f'/share/{share["token"]}/download').content == b"original bytes"
    post(f"/drive/folder/{root}/trash", {})
    assert len(c.get("/drive/items?view=trash", headers=headers).json()["items"]) == 3
    assert c.get(f'/share/{share["token"]}/download').status_code == 403
    post(f"/drive/folder/{root}/restore", {})
    assert c.get("/drive/items?view=trash", headers=headers).json()["items"] == []
    post(f"/drive/file/{file_id}/move", {"parent_id": None})
    assert any(i["id"] == file_id and i["kind"] == "file" for i in c.get("/drive/items", headers=headers).json()["items"])
    assert c.post(f"/drive/file/{file_id}/delete", headers=headers, json={}).status_code == 400
    post(f"/drive/file/{file_id}/trash", {})
    post(f"/drive/file/{file_id}/delete", {})
    assert c.get(f"/files/{file_id}/download", headers=headers).status_code == 404


def test_versions_and_independent_trash(local_client):
    c = local_client
    headers = account(c, "versions@example.com")
    root = c.post("/drive/folders", headers=headers, json={"name": "Versions"}).json()["id"]
    file_id = c.post("/files/upload", headers=headers, data={"folder_id": root},
                     files={"file": ("original.txt", b"first", "text/plain")}).json()["id"]
    replaced = c.put(f"/files/{file_id}/replace", headers=headers,
                     files={"file": ("updated.txt", b"second", "text/plain")})
    assert replaced.status_code == 200, replaced.text
    versions = c.get(f"/files/{file_id}/versions", headers=headers).json()
    assert len(versions) == 1
    assert c.get(f"/files/{file_id}/download", headers=headers).content == b"second"
    restored = c.post(f'/files/{file_id}/versions/{versions[0]["id"]}/restore', headers=headers)
    assert restored.status_code == 200, restored.text
    assert c.get(f"/files/{file_id}/download", headers=headers).content == b"first"
    assert c.post(f"/drive/file/{file_id}/trash", headers=headers, json={}).status_code == 200
    assert c.post(f"/drive/folder/{root}/trash", headers=headers, json={}).status_code == 200
    assert c.post(f"/drive/folder/{root}/restore", headers=headers, json={}).status_code == 200
    trash = c.get("/drive/items?view=trash", headers=headers).json()["items"]
    assert len(trash) == 1 and trash[0]["id"] == file_id


def test_account_profile_settings_password_and_support(local_client):
    import base64
    c = local_client
    headers = account(c, "profile@example.com")
    other = account(c, "other.profile@example.com")
    original_token = dict(headers)
    settings = {"theme": "dark", "density": "compact", "default_view": "drive", "layout": "grid"}
    r = c.put("/account/settings", headers=headers, json=settings)
    assert r.status_code == 200, r.text
    assert c.get("/me", headers=headers).json()["settings"] == settings
    assert c.get("/me", headers=other).json()["settings"]["theme"] == "light"
    png = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+a7X0AAAAASUVORK5CYII=")
    assert c.post("/account/avatar", headers=headers, files={"file": ("avatar.svg", b"<svg></svg>", "image/svg+xml")}).status_code == 400
    r = c.post("/account/avatar", headers=headers, files={"file": ("avatar.png", png, "image/png")})
    assert r.status_code == 200 and r.json()["avatar"].startswith("data:image/png;base64,")
    assert c.get("/me", headers=other).json()["avatar"] is None
    assert c.delete("/account/avatar", headers=headers).json()["avatar"] is None
    assert c.patch("/account", headers=headers, json={"full_name": "Updated", "email": "new@example.com"}).status_code == 400
    r = c.patch("/account", headers=headers, json={"full_name": "Updated", "email": "new@example.com", "current_password": "strong-test-pass"})
    assert r.status_code == 200, r.text
    headers = {"Authorization": "Bearer " + r.json()["access_token"]}
    assert c.get("/me", headers=original_token).status_code == 401
    assert c.get("/me", headers=headers).json()["full_name"] == "Updated"
    assert c.post("/account/password", headers=headers, json={"current_password": "wrong", "new_password": "replacement-password"}).status_code == 400
    r = c.post("/account/password", headers=headers, json={"current_password": "strong-test-pass", "new_password": "replacement-password"})
    assert r.status_code == 200, r.text
    assert c.get("/me", headers=headers).status_code == 401
    headers = {"Authorization": "Bearer " + r.json()["access_token"]}
    assert c.get("/me", headers=headers).status_code == 200
    assert c.post("/auth/login", data={"username": "new@example.com", "password": "strong-test-pass"}).status_code == 401
    assert c.post("/auth/login", data={"username": "new@example.com", "password": "replacement-password"}).status_code == 200
    r = c.post("/support", headers=headers, json={"category": "upload", "subject": "Test issue", "description": "Disposable test report for integration checks."})
    assert r.status_code == 201, r.text
    assert c.get("/support", headers=headers).json()[0]["id"] == r.json()["id"]
    assert c.get("/support", headers=other).json() == []


def test_advanced_search_content_location_type_and_dates(local_client):
    from datetime import date, timedelta
    c = local_client
    headers = account(c, "search.advanced@example.com")
    other = account(c, "search.other@example.com")
    folder = c.post("/drive/folders", headers=headers, json={"name": "Research"}).json()["id"]
    file_id = c.post("/files/upload", headers=headers, data={"folder_id": folder},
                     files={"file": ("notes.md", b"Project amber details and milestones", "text/markdown")}).json()["id"]
    c.post("/files/upload", headers=headers, files={"file": ("outside.txt", b"Other words", "text/plain")})
    def search(**params):
        return c.get("/drive/items", headers=headers, params={"advanced": True, **params})
    r = search(content_q="AMBER milestones", file_type="documents", location_id=folder)
    assert r.status_code == 200, r.text
    assert len(r.json()["items"]) == 1 and r.json()["items"][0]["id"] == file_id
    assert search(file_type="videos").json()["items"] == []
    assert search(file_type="folders").json()["items"][0]["name"] == "Research"
    assert search(modified_after=str(date.today() + timedelta(days=1))).json()["items"] == []
    assert search(modified_after="2030-01-01", modified_before="2020-01-01").status_code == 400
    assert c.get("/drive/items", headers=other, params={"advanced": True, "location_id": folder}).status_code == 404
    c.post(f"/drive/file/{file_id}/star", headers=headers, json={"starred": True})
    assert len(search(starred_only=True).json()["items"]) == 1
    c.post(f"/drive/file/{file_id}/trash", headers=headers, json={})
    assert search(content_q="amber").json()["items"] == []
    assert len(search(view="trash", content_q="amber").json()["items"]) == 1
